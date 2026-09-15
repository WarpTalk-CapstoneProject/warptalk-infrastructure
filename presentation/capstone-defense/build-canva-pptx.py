"""Builds a PPTX whose text is real PowerPoint text, so Canva turns it into editable
elements on import. Anything that cannot survive as a native shape — logos, the SVG
illustrations, the architecture diagram, the gradient/noise cards' texture — is placed
as a picture instead."""
import json, os, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

DIR = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(DIR, '.pptx-native')
EMU_PX = 6350                       # 12192000 EMU / 1920 px
FONT = 'Avenir Next'
GRADIENT_INK = RGBColor(0x5B, 0x56, 0xFF)   # gradient-clipped text falls back to its mid stop

def emu(px): return Emu(int(round(px * EMU_PX)))
def rgb(css):
    m = re.findall(r'[\d.]+', css or '')
    if len(m) < 3: return None
    if len(m) > 3 and float(m[3]) < 0.06: return None       # effectively transparent
    return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))

ALIGN = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT}

deck = json.load(open(os.path.join(SRC, 'layout.json')))
by_name = {d['name']: d for d in deck}

# full 30-slide order; covers ship as images because their gradient/noise/wave layers
# have no native equivalent, and they are already in Canva from the earlier upload
ORDER = [
 ('cover', 'warptalk-cover-editorial-v2.png'), ('cover', 'slide-02-team.png'), ('cover', 'slide-03-contents.png'),
 ('cover', 'section-01-context.png'),      ('native', '05-context'), ('native', '06-problems'),
 ('cover', 'section-02-solution.png'),     ('native', '08-solution'),
 ('cover', 'section-03-architecture.png'), ('native', '10-architecture'),
 ('cover', 'section-04-tech.png'),         ('native', '12-tech-stack'),
 ('cover', 'section-05-requirements.png'), ('native', '14-actors'), ('native', '15-participants'),
 ('native', '16-host'), ('native', '17-workspace-owner'), ('native', '18-administrator'),
 ('cover', 'section-06-workflows.png'),    ('native', '20-flow-setup'), ('native', '21-flow-meeting'),
 ('native', '22-flow-pipeline'), ('native', '23-flow-record'),
 ('cover', 'section-07-achievements.png'), ('native', '25-limitations'), ('native', '26-achievements'),
 ('native', '27-future-work'),
 ('cover', 'section-08-qa.png'),           ('native', '29-thank-you'), ('native', '30-qa'),
]

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]

stats = {'text': 0, 'run': 0, 'shape': 0, 'pic': 0, 'coverimg': 0}

for kind, ref in ORDER:
    slide = prs.slides.add_slide(blank)
    if kind == 'cover':
        slide.shapes.add_picture(os.path.join(DIR, ref), 0, 0, prs.slide_width, prs.slide_height)
        stats['coverimg'] += 1
        continue

    d = by_name[ref]
    for s in d['surfaces']:                                     # backgrounds and rules first
        if s['borderTop'] or s['borderBottom']:
            for edge, col in (('borderTop', s['borderTop']), ('borderBottom', s['borderBottom'])):
                c = rgb(col)
                if not c: continue
                y = s['y'] if edge == 'borderTop' else s['y'] + s['h'] - 1
                r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(s['x']), emu(y), emu(s['w']), emu(1))
                r.fill.solid(); r.fill.fore_color.rgb = c; r.line.fill.background(); r.shadow.inherit = False
                stats['shape'] += 1
        c = rgb(s['fill'])
        if not c and not s['gradient']: continue
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if s['radius'] >= 6 else MSO_SHAPE.RECTANGLE
        sh = slide.shapes.add_shape(shape, emu(s['x']), emu(s['y']), emu(s['w']), emu(s['h']))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE and s['w']:
            sh.adjustments[0] = min(0.5, s['radius'] / min(s['w'], s['h']))
        if s['gradient']:
            sh.fill.gradient(); sh.fill.gradient_angle = 140
            sh.fill.gradient_stops[0].color.rgb = RGBColor(0x65, 0x76, 0xEE)
            sh.fill.gradient_stops[1].color.rgb = RGBColor(0x35, 0x45, 0xE7)
        else:
            sh.fill.solid(); sh.fill.fore_color.rgb = c
        sh.line.fill.background(); sh.shadow.inherit = False
        stats['shape'] += 1

    for p in d['pictures']:                                     # then the things that must stay pixels
        f = os.path.join(SRC, p.get('file', ''))
        if p.get('file') and os.path.exists(f):
            slide.shapes.add_picture(f, emu(p['x']), emu(p['y']), emu(p['w']), emu(p['h']))
            stats['pic'] += 1

    for t in d['texts']:                                        # text last, so it sits on top
        box = slide.shapes.add_textbox(emu(t['x'] - 2), emu(t['y'] - 2), emu(t['w'] + 8), emu(t['h'] + 8))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        para = tf.paragraphs[0]
        para.alignment = ALIGN.get(t['align'], PP_ALIGN.LEFT)
        if t['lineHeight']: para.line_spacing = Pt(t['lineHeight'] * 0.75)
        for r in t['runs']:                                  # inline runs keep their own weight and colour
            if r['text'] == '\n':
                para = tf.add_paragraph()
                para.alignment = ALIGN.get(t['align'], PP_ALIGN.LEFT)
                if t['lineHeight']: para.line_spacing = Pt(t['lineHeight'] * 0.75)
                continue
            run = para.add_run()
            run.text = r['text'].upper() if r['upper'] else r['text']
            run.font.name = FONT
            run.font.size = Pt(round(r['size'] * 0.75, 1))
            w = r['weight']
            run.font.bold = int(w) >= 600 if w.isdigit() else w == 'bold'
            run.font.color.rgb = GRADIENT_INK if r['color'] == 'GRADIENT' else (rgb(r['color']) or RGBColor(0x11, 0x11, 0x13))
            stats['run'] += 1
        stats['text'] += 1

out = os.path.join(DIR, 'canva-upload', 'warptalk-capstone-30-slides-editable.pptx')
os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print('slides:', len(prs.slides.__iter__.__self__._sldIdLst))
print('native text boxes: %(text)d (%(run)d runs) | native shapes: %(shape)d | pictures: %(pic)d | cover images: %(coverimg)d' % stats)
print(out, '%.1f MB' % (os.path.getsize(out) / 1e6))
